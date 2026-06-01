"""finalize_run_v2 — shared tail for dry-run + draft.

Both flows end the same way: web-enhance opt-in values → heal dates →
heal user-input prose → substitute placeholders into the template.docx
→ upload filled docx to R2 → return presigned URL + warnings. This
module owns the shared tail; the dry-run / draft routers call
`finalize_run_v2(...)` and wrap its `FinalizedRunV2` into the
route-specific response type.

Load-bearing ORDERING INVARIANT: web-enhance BEFORE date-heal BEFORE
prose-heal BEFORE fill. Web enhance runs first so its output flows
through date + prose heal (web-search-derived dates get normalized;
prose-heal grammars the enhanced value into the paragraph). Date heal
locks dates to the firm default. Prose-heal's LLM prompt sees
already-normalized strings (and is instructed not to touch them).
Fill needs healed values so the rendered docx is polished.

R2 layout (per plan):
    cases/{case_id}/{output_prefix}/{uuid}.docx

where `output_prefix` is "dry_run" (ephemeral) or "draft/{task_id}"
(persistent — Phase 3 drafting). Slice C ships the dry-run path
only; Phase 3 wires the draft path with its own output_prefix scheme.
"""

from __future__ import annotations

import logging
import uuid
from io import BytesIO
from typing import Any

from docx import Document

from src.core.common.documents.docx_template import DocxTemplateService
from src.core.common.storage.r2 import r2_service

from ..agents.heal import UserInputHealAgentV2
from ..agents.web_enhance import WebEnhanceAgentV2
from ..resolvers.date_healing import DateHealingResolverV2
from ..types.fields import TemplateFieldV2, TemplateSpecV2
from ..types.orchestration import FinalizedRunV2, GrammarRepairV2
from ..types.resolution import ResolvedTemplateValueV2

logger = logging.getLogger(__name__)


async def finalize_run_v2(
    *,
    template_id: str,
    case_id: str,
    spec: TemplateSpecV2,
    all_resolved: list[ResolvedTemplateValueV2],
    template_bytes: bytes | None = None,
    output_prefix: str = "dry_run",
    resource_key: str | None = None,
) -> FinalizedRunV2:
    """Heal → fill → upload → presigned URL.

    Args:
        template_id: The templates_v2 row id; used to download the
            placeholder-marked template.docx from R2 when `template_bytes`
            isn't supplied.
        case_id: The case to render against (drives the R2 upload path).
        spec: The TemplateSpecV2 — needed by UserInputHealAgentV2 to look
            up per-field params (heal target, output_expectation).
        all_resolved: Every resolved value from `run_initial_stages_v2`
            (no pause) or `run_resume_stages_v2` (post-pause).
        template_bytes: Pre-downloaded template.docx. When None, this
            function downloads it from R2.
        output_prefix: R2 key prefix within the case's folder
            ("dry_run" by default; Phase 3 draft path overrides).
        resource_key: R2 folder name for the upload. Defaults to
            `case_id` — Phase 1 unfiled-cases pattern (legacy
            sanitized-slug cases pass their `legacy_id` instead).
    """
    rk = resource_key or case_id
    if template_bytes is None:
        template_bytes = await r2_service.download_file(
            template_id=template_id,
            filename="template.docx",
            prefix="template_v2",
        )

    # 1. Web enhance — opt-in per field (`params.web_enhance_instruction`).
    #    Runs first so date-heal can normalize any dates the search
    #    pulls in and prose-heal can shape the enhanced value into the
    #    paragraph. Soft-fail: unenhanced value passes through on any
    #    error.
    web_enhanced = await _apply_web_enhance(
        spec=spec,
        resolved_values=all_resolved,
        template_bytes=template_bytes,
    )

    # 2. Date heal — deterministic, runs after enhancement so any
    #    enhancement-produced dates get normalized too.
    date_healed = DateHealingResolverV2.apply(web_enhanced)

    # 3. Prose heal — LLM per user-input field; uses template_bytes for
    #    surrounding-paragraph awareness. Soft failure: returns raw
    #    values on LLM error.
    prose_healed = await UserInputHealAgentV2.heal_resolved_values(
        template_bytes=template_bytes,
        template_fields=spec.fields,
        resolved_values=date_healed,
    )

    # 4. Fill docx — substitute every [[var]] with the resolved value.
    filled_bytes, unresolved = _fill_template_v2(
        template_bytes=template_bytes,
        spec=spec,
        resolved_values=prose_healed,
    )

    # 4b. Tier 2 LLM-assisted format autofix (behind
    #     TEMPLATE_FORMAT_AUTOFIX_V2 flag). Detects drift introduced
    #     by `_substitute_placeholder` (e.g. multi-line resolved
    #     values inline-joined with " and " when they should have
    #     rendered with `<w:br/>`), calls Sonnet 4.6 per drifted
    #     paragraph, content-equality-checks the response, rebuilds
    #     the XML deterministically. Soft-fails to unfixed bytes if
    #     any step errors. Awaited natively (NOT thread-bridged) so
    #     the LangChain httpx connection pool stays bound to the
    #     same event loop we're running in.
    resolved_value_map = {
        rv.template_variable: (rv.value or "").strip()
        for rv in prose_healed
        if rv.value and rv.value.strip()
    }
    filled_bytes = await DocxTemplateService.maybe_autofix_fill_async(
        template_bytes=template_bytes,
        filled_bytes=filled_bytes,
        resolved_values=resolved_value_map,
    )

    # 4c. Tier 2 LLM-assisted grammar autofix (behind
    #     TEMPLATE_GRAMMAR_AUTOFIX_V2 flag). Catches plural/singular
    #     agreement mismatches between hardcoded template language
    #     ("Debtors", "their", "have") and the resolved values'
    #     actual cardinality — e.g. single-debtor cases where the
    #     template was authored against joint debtors. Runs AFTER
    #     the format autofix so layout drift is gone first; the
    #     grammar fixer sees the cleaned bytes and only addresses
    #     word-level agreement. Soft-fails to unfixed bytes on any
    #     error. Captures the list of applied swaps so the FE
    #     Resolution Log can show forensic detail.
    filled_bytes, grammar_repair_records = (
        await DocxTemplateService.maybe_autofix_grammar_async(
            filled_bytes=filled_bytes,
            resolved_values=resolved_value_map,
        )
    )
    grammar_repairs = [
        GrammarRepairV2(
            paragraph_index=r.paragraph_index,
            original_word=r.original_word,
            replacement_word=r.replacement_word,
            occurrences=r.occurrences,
            paragraph_preview=r.paragraph_preview,
            reason=r.reason,
        )
        for r in grammar_repair_records
    ]

    # 5. Upload to R2 + return a presigned URL.
    generated_filename = f"{output_prefix}/{uuid.uuid4()}.docx"
    await r2_service.upload_file(
        file_content=filled_bytes,
        template_id=rk,
        filename=generated_filename,
        prefix="cases",
    )
    generated_doc_url = await r2_service.get_presigned_url(
        template_id=rk,
        filename=generated_filename,
        prefix="cases",
    )
    r2_object_key = f"cases/{rk}/{generated_filename}"

    return FinalizedRunV2(
        resolved_values=prose_healed,
        generated_doc_url=generated_doc_url,
        r2_object_key=r2_object_key,
        unresolved=unresolved,
        warnings=_build_warnings(prose_healed, unresolved),
        grammar_repairs=grammar_repairs,
        filled_bytes=filled_bytes,
    )


# ─── web enhance pass ────────────────────────────────────────────────


async def _apply_web_enhance(
    *,
    spec: TemplateSpecV2,
    resolved_values: list[ResolvedTemplateValueV2],
    template_bytes: bytes,
) -> list[ResolvedTemplateValueV2]:
    """For every resolved row whose field has a non-empty
    `web_enhance_instruction`, run `WebEnhanceAgentV2.run(...)` and
    swap the resolved value with the enhanced output.

    Soft-fail at every level: missing field, empty current_value,
    empty instruction, agent error — all degrade to passing the
    original row through unchanged.

    Sequential (not parallel) — web search is rate-limited at the
    Anthropic side and these are rarely more than 1-2 fields per
    template. Can be parallelized with `asyncio.gather` later if it
    matters.
    """
    fields_by_name: dict[str, TemplateFieldV2] = {
        f.template_variable: f for f in spec.fields
    }
    out: list[ResolvedTemplateValueV2] = []
    for rv in resolved_values:
        field = fields_by_name.get(rv.template_variable)
        instruction = (
            (field.params.web_enhance_instruction or "").strip()
            if field and field.params is not None
            else ""
        )
        if not instruction or not rv.value or not rv.value.strip():
            out.append(rv)
            continue
        paragraph = DocxTemplateService.find_paragraph_containing(
            template_bytes, f"[[{rv.template_variable}]]",
        )
        output_expectation = (
            field.params.output_expectation if field and field.params else None
        )
        enhanced = await WebEnhanceAgentV2.run(
            variable_name=rv.template_variable,
            current_value=rv.value,
            web_enhance_instruction=instruction,
            template_property_marker=field.template_property_marker if field else None,
            template_paragraph=paragraph,
            case_details=None,
            output_expectation=output_expectation,
        )
        if enhanced == rv.value:
            out.append(rv)
            continue
        out.append(rv.model_copy(update={
            "value": enhanced,
            "note": _append_note(rv.note, "web-enhanced"),
        }))
    return out


def _append_note(existing: str, addition: str) -> str:
    if not existing:
        return addition
    return f"{existing}; {addition}"


# ─── docx fill (v2-native; reuses v1's substitution primitive) ───────


def _fill_template_v2(
    *,
    template_bytes: bytes,
    spec: TemplateSpecV2,
    resolved_values: list[ResolvedTemplateValueV2],
) -> tuple[bytes, list[str]]:
    """Substitute every `[[template_variable]]` placeholder with its
    resolved value. Returns `(filled_bytes, unresolved_placeholders)`.

    Mirrors v1's `DocxTemplateService.fill_template` semantically but
    reads v2's `TemplateFieldV2` shape directly (no v1 TemplateField
    adapter). The substitution primitive
    (`DocxTemplateService._substitute_placeholder`) is reused as a
    pure utility.
    """
    by_name = {rv.template_variable: rv for rv in resolved_values}
    unresolved: list[str] = []
    resolved_value_map: dict[str, str] = {}

    doc = Document(BytesIO(template_bytes))
    # Snapshot per-paragraph caption-shape decisions BEFORE any
    # substitution mutates the runs. See `fill_template` for the
    # rationale — without this, multi-placeholder paragraphs misroute
    # to inline " and " on the second placeholder onwards.
    caption_shape_map = DocxTemplateService._compute_caption_shape_map(doc)
    for field in spec.fields:
        placeholder = f"[[{field.template_variable}]]"
        rv = by_name.get(field.template_variable)
        value = (rv.value if rv else "").strip()
        if value:
            DocxTemplateService._substitute_placeholder(
                doc, placeholder, value, caption_shape_map=caption_shape_map,
            )
            resolved_value_map[field.template_variable] = value
        else:
            unresolved.append(placeholder)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    filled_bytes = out.read()

    # Tier 1 + Tier 2 format-drift gate (validator + LLM fixer behind
    # TEMPLATE_FORMAT_AUTOFIX_V2 env flag). Runs identical to the v1
    # `fill_template`'s tail — same static helper, same env flag.
    filled_bytes = DocxTemplateService._maybe_run_fill_validator(
        template_bytes, filled_bytes, resolved_value_map,
    )
    return filled_bytes, unresolved


def _build_warnings(
    resolved_values: list[ResolvedTemplateValueV2],
    unresolved: list[str],
) -> list[str]:
    """Surface unresolved placeholders and low-confidence extractions
    as human-readable warnings. The dry-run result modal shows these."""
    warnings: list[str] = []
    for placeholder in unresolved:
        warnings.append(f"Unresolved placeholder: {placeholder}")
    for rv in resolved_values:
        if rv.confidence in ("low", "none"):
            warnings.append(
                f"{rv.confidence.title()}-confidence value for "
                f"'{rv.template_variable}': {rv.note or '(no note)'}"
            )
    return warnings
